# -*- coding: utf-8 -*-
"""
RIFAN — Services Presentation Generator
هوية بصرية مطابقة للموقع: عاجي / دكتور / ذهبي · El Messiri / Tajawal / Marcellus
"""
import os
import qrcode
from qrcode.constants import ERROR_CORRECT_Q
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from io import BytesIO

# ---------------------------------------------------------------- palette
IVORY      = RGBColor(0xFB, 0xF9, 0xF6)
IVORY_DEEP = RGBColor(0xF4, 0xF0, 0xE9)
SAND       = RGBColor(0xEC, 0xE6, 0xDA)
LINE       = RGBColor(0xE6, 0xDF, 0xD2)
INK        = RGBColor(0x21, 0x1D, 0x18)
INK_SOFT   = RGBColor(0x4A, 0x44, 0x3C)
MUTED      = RGBColor(0x87, 0x7E, 0x70)
GOLD       = RGBColor(0xB0, 0x8D, 0x4F)
GOLD_DEEP  = RGBColor(0x96, 0x76, 0x3A)
GOLD_SOFT  = RGBColor(0xD9, 0xC9, 0xA8)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

F_HEAD = "El Messiri"   # عناوين فاخرة
F_BODY = "Tajawal"      # نصوص عربية
F_LAT  = "Marcellus"    # لاتيني

SW, SH = Inches(13.333), Inches(7.5)
LOGO = r"C:\Projects\rivan\logo_rifan_new.png"
IMG = r"C:\Projects\rivan\pptx_img"
WHATSAPP = "966500939971"  # رقم الواتساب بصيغة دولية بدون +

# ---------------------------------------------------------------- helpers
def set_font(run, name, size, color, bold=False, spc=None):
    f = run.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    # complex-script font (Arabic)
    cs = rPr.find(qn('a:cs'))
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {})
        rPr.append(cs)
    cs.set('typeface', name)
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)
    if spc:
        rPr.set('spc', str(spc))

def para(tf, text, font=F_BODY, size=14, color=INK, bold=False, align=PP_ALIGN.RIGHT,
         rtl=True, spc=None, space_after=0, line=1.0, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    if line != 1.0:
        p.line_spacing = line
    if space_after:
        p.space_after = Pt(space_after)
    if rtl:
        pPr = p._p.get_or_add_pPr()
        pPr.set('rtl', '1')
    r = p.add_run()
    r.text = text
    set_font(r, font, size, color, bold, spc)
    return p

def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf

def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None, shadow=False):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if not shadow:
        # explicit no-shadow
        spPr = sp._element.spPr
        el = spPr.makeelement(qn('a:effectLst'), {})
        spPr.append(el)
    return sp

def hline(slide, x, y, w, color=GOLD, weight=1.2):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(weight))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background()
    ln.shadow.inherit = False
    return ln

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def latin_label(slide, x, y, w, text, color=MUTED, size=11, align=PP_ALIGN.RIGHT):
    _, tf = box(slide, x, y, w, Inches(0.3))
    para(tf, text, F_LAT, size, color, align=align, rtl=False, spc=350, first=True)

def corner_marks(slide, x, y, w, h, color=GOLD, arm=Inches(0.55), t=Pt(1.5)):
    # زوايا ذهبية (إطار فاخر بأربع زوايا L)
    for (cx, cy, dx, dy) in [
        (x, y, 1, 1), (x + w, y, -1, 1),
        (x, y + h, 1, -1), (x + w, y + h, -1, -1),
    ]:
        hln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
            cx if dx == 1 else cx - arm, cy - t / 2 if False else cy, arm if dx == 1 else arm, t)
        hln.fill.solid(); hln.fill.fore_color.rgb = color
        hln.line.fill.background(); hln.shadow.inherit = False
        vln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
            cx - t / 2 if dx == 1 else cx - t / 2, cy if dy == 1 else cy - arm, t, arm)
        vln.fill.solid(); vln.fill.fore_color.rgb = color
        vln.line.fill.background(); vln.shadow.inherit = False

from PIL import Image as PILImage

def crop(src_name, out_name, ratio):
    """قص مركزي بنسبة عرض/ارتفاع محددة لتجنّب أي تشويه"""
    src = os.path.join(IMG, src_name)
    out = os.path.join(IMG, out_name)
    im = PILImage.open(src).convert("RGB")
    W, H = im.size
    if W / H > ratio:
        nw = int(H * ratio); x0 = (W - nw) // 2; box = (x0, 0, x0 + nw, H)
    else:
        nh = int(W / ratio); y0 = (H - nh) // 2; box = (0, y0, W, y0 + nh)
    im.crop(box).save(out, quality=88)
    return out

def overlay(slide, x, y, w, h, color=INK, opacity=80):
    """طبقة لون شبه شفافة فوق الصور"""
    sp = rect(slide, x, y, w, h, fill=color)
    solid = sp._element.spPr.find(qn('a:solidFill'))
    srgb = solid.find(qn('a:srgbClr'))
    srgb.append(srgb.makeelement(qn('a:alpha'), {'val': str(int(opacity * 1000))}))
    return sp

def make_qr(url, path, dark="#211d18", light="#fbf9f6"):
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_Q, box_size=12, border=2)
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color=dark, back_color=light)
    img.save(path)

def slide_header(slide, label, title):
    latin_label(slide, Inches(0.7), Inches(0.52), Inches(12.0), label, GOLD_DEEP, 11)
    _, tf = box(slide, Inches(0.7), Inches(0.82), Inches(12.0), Inches(0.85))
    para(tf, title, F_HEAD, 32, INK, bold=True, first=True)
    hline(slide, Inches(11.1), Inches(1.62), Inches(1.6), GOLD, 2.0)
    # شعار مائي لاتيني يسار أسفل
    _, tf2 = box(slide, Inches(0.7), Inches(7.05), Inches(3), Inches(0.3))
    para(tf2, "RIFAN", F_LAT, 10, LINE, align=PP_ALIGN.LEFT, rtl=False, spc=500, first=True)

def gold_num(slide, x, y, n):
    c = rect(slide, x, y, Inches(0.52), Inches(0.52), fill=None, line_color=GOLD, line_w=1.1,
             shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = n
    set_font(r, F_LAT, 15, GOLD_DEEP, spc=0)

# ---------------------------------------------------------------- build
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

def new_slide(color=IVORY):
    s = prs.slides.add_slide(BLANK)
    bg(s, color)
    return s

# ================================================================ 1 · الغلاف
s = new_slide(INK)
cover = crop("hero.jpg", "hero_169.jpg", 13.333 / 7.5)
s.shapes.add_picture(cover, 0, 0, SW, SH)
overlay(s, 0, 0, SW, SH, INK, 82)
rect(s, Inches(0.32), Inches(0.32), SW - Inches(0.64), SH - Inches(0.64),
     fill=None, line_color=GOLD, line_w=1.0)
corner_marks(s, Inches(0.55), Inches(0.55), SW - Inches(1.1), SH - Inches(1.1), GOLD, Inches(0.7), Pt(2.0))
s.shapes.add_picture(LOGO, SW/2 - Inches(1.0), Inches(1.15), height=Inches(2.0))
_, tf = box(s, Inches(1.5), Inches(3.35), Inches(10.333), Inches(1.1))
para(tf, "رِفان", F_HEAD, 60, GOLD_SOFT, bold=True, align=PP_ALIGN.CENTER, first=True)
_, tf = box(s, Inches(1.5), Inches(4.45), Inches(10.333), Inches(0.6))
para(tf, "إدارة وتشغيل الوحدات السكنية للإيجار قصير المدى", F_BODY, 22, IVORY,
     align=PP_ALIGN.CENTER, first=True)
hline(s, SW/2 - Inches(1.1), Inches(5.35), Inches(2.2), GOLD, 1.5)
_, tf = box(s, Inches(1.5), Inches(5.55), Inches(10.333), Inches(0.4))
para(tf, "PROFESSIONAL SHORT-TERM RENTAL MANAGEMENT", F_LAT, 12, GOLD_SOFT,
     align=PP_ALIGN.CENTER, rtl=False, spc=400, first=True)
_, tf = box(s, Inches(1.5), Inches(6.45), Inches(10.333), Inches(0.4))
para(tf, "عرض تقديمي للخدمات", F_BODY, 12, GOLD_SOFT,
     align=PP_ALIGN.CENTER, first=True)

# ================================================================ 2 · المحتويات
s = new_slide()
slide_header(s, "AGENDA", "محتويات العرض")
items = ["من نحن", "رحلة المالك", "خدماتنا", "آلية العمل",
         "لماذا رِفان", "منصات الحجز", "قصص نجاح", "الأسئلة الشائعة والتواصل"]
cw = Inches(5.8)
xs = [SW - Inches(0.7) - cw, Inches(0.7)]  # RTL: العمود الأيمن أولًا
for i, it in enumerate(items):
    col, row = i // 4, i % 4
    x = xs[col]
    y = Inches(2.15) + row * Inches(1.15)
    card = rect(s, x, y, cw, Inches(0.95),
                fill=WHITE, line_color=LINE, line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    gold_num(s, x + cw - Inches(0.75), y + Inches(0.215), f"0{i+1}")
    _, tf = box(s, x + Inches(0.3), y + Inches(0.26), cw - Inches(1.15), Inches(0.45))
    para(tf, it, F_HEAD, 17, INK, bold=True, first=True)

# ================================================================ 3 · من نحن
s = new_slide()
slide_header(s, "ABOUT RIFAN", "رِفان... أكثر من إدارة وحدة")
_, tf = box(s, Inches(4.6), Inches(2.15), Inches(8.1), Inches(2.6))
para(tf, "في رِفان، نؤمن أن نجاح الوحدة السكنية لا يعتمد على عرضها فقط، بل على جودة إدارتها اليومية والتفاصيل التي تصنع تجربة الضيف.",
     F_BODY, 16, INK_SOFT, line=1.6, space_after=14, first=True)
para(tf, "نتولى إدارة وتشغيل الوحدة باحترافية، من إدارة الحجوزات وخدمة الضيوف إلى التنسيق مع فرق التنظيف والصيانة، ومتابعة الأداء والتقييمات.",
     F_BODY, 16, INK_SOFT, line=1.6)
hline(s, Inches(12.15), Inches(5.0), Inches(0.55), GOLD, 2.0)
_, tf = box(s, Inches(4.6), Inches(5.15), Inches(8.1), Inches(0.4))
para(tf, "Professional Short-Term Rental Management", F_LAT, 12, MUTED, rtl=False, spc=300, first=True)
# لوحة الصور (بنفس أسلوب الموقع): صورة طولية بإطار ذهبي مزاح + صورة صغيرة متراكبة + شارة
tall = crop("about.jpg", "about_tall.jpg", 3.5 / 4.35)
small = crop("about_small.jpg", "about_small43.jpg", 4 / 3)
rect(s, Inches(1.15), Inches(1.95), Inches(3.5), Inches(4.35), fill=None,
     line_color=GOLD, line_w=1.2)  # إطار ذهبي مزاح
s.shapes.add_picture(tall, Inches(0.9), Inches(2.15), Inches(3.5), Inches(4.35))
s.shapes.add_picture(small, Inches(3.05), Inches(5.1), Inches(1.9), Inches(1.425))
rect(s, Inches(3.05), Inches(5.1), Inches(1.9), Inches(1.425), fill=None, line_color=IVORY, line_w=3.5)
badge = rect(s, Inches(1.35), Inches(6.1), Inches(1.5), Inches(0.5), fill=INK,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
tfb = badge.text_frame; tfb.margin_left = tfb.margin_right = 0
pb = tfb.paragraphs[0]; pb.alignment = PP_ALIGN.CENTER
rb = pb.add_run(); rb.text = "RIFAN"
set_font(rb, F_LAT, 12, GOLD, spc=300)

# ================================================================ 3.5 · معرض الوحدات
s = new_slide()
slide_header(s, "OUR UNITS", "لمحة من وحدات تديرها رِفان")
gallery = [
    ("bedroom.jpg", "غرف نوم بلمسة ضيافة فاخرة"),
    ("living.jpg", "مساحات معيشة مريحة وأنيقة"),
    ("cinematic.jpg", "تفاصيل مصممة لتجربة ضيف متكاملة"),
]
gw = Inches(3.95)
for i, (img, cap) in enumerate(gallery):
    x = Inches(12.65) - gw - i * Inches(4.12)
    y = Inches(2.2)
    g = crop(img, f"g{i}.jpg", 3.95 / 2.9)
    s.shapes.add_picture(g, x, y, gw, Inches(2.9))
    rect(s, x, y, gw, Inches(2.9), fill=None, line_color=GOLD, line_w=1.0)
    _, tf = box(s, x, y + Inches(3.05), gw, Inches(0.9))
    para(tf, cap, F_BODY, 13.5, INK_SOFT, align=PP_ALIGN.CENTER, line=1.35, first=True)
_, tf = box(s, Inches(0.7), Inches(6.35), Inches(12.0), Inches(0.5))
para(tf, "التجهيز والتصوير الاحترافي جزء من خطة التشغيل التي نبدأ بها كل وحدة جديدة.",
     F_BODY, 13.5, GOLD_DEEP, align=PP_ALIGN.CENTER, first=True)

# ================================================================ 4 · رحلة المالك
s = new_slide()
slide_header(s, "OWNER JOURNEY", "رحلة المالك مع رِفان")
journey = [("أنت", "تملك الوحدة", True), ("رِفان", "تتولى الإدارة والتشغيل", False),
           ("الضيف", "يحصل على تجربة متكاملة", False), ("أنت", "تتابع أداء وحدتك بوضوح", True)]
w = Inches(2.85); gap = Inches(0.22)
for i, (actor, text, you) in enumerate(journey):
    x = Inches(12.65) - w - i * (w + gap)
    y = Inches(2.6)
    card = rect(s, x, y, w, Inches(2.9),
                fill=INK if you else WHITE,
                line_color=None if you else LINE, line_w=0.75,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    gold_num(s, x + w/2 - Inches(0.26), y + Inches(0.35), f"0{i+1}")
    _, tf = box(s, x + Inches(0.25), y + Inches(1.1), w - Inches(0.5), Inches(0.5))
    para(tf, actor, F_HEAD, 20, GOLD if you else GOLD_DEEP, bold=True, align=PP_ALIGN.CENTER, first=True)
    _, tf = box(s, x + Inches(0.3), y + Inches(1.7), w - Inches(0.6), Inches(1.0))
    para(tf, text, F_BODY, 14, IVORY if you else INK_SOFT, align=PP_ALIGN.CENTER, line=1.4, first=True)
_, tf = box(s, Inches(0.7), Inches(6.0), Inches(12.0), Inches(0.5))
para(tf, "الوحدة تبقى ملكك بالكامل — ورِفان شريك التشغيل الذي يحوّلها إلى تجربة ضيافة رابحة.",
     F_BODY, 14, GOLD_DEEP, align=PP_ALIGN.CENTER, first=True)

# ================================================================ 5-6 · الخدمات
services = [
    ("إدارة الحجوزات", "متابعة الحجوزات وتنظيمها وإدارة تفاصيل الإقامة."),
    ("خدمة الضيوف", "التواصل مع الضيوف ومتابعة احتياجاتهم قبل وأثناء وبعد الإقامة."),
    ("التشغيل اليومي", "إدارة التفاصيل التشغيلية للوحدة ومتابعة سير العمل."),
    ("التنظيف والتجهيز", "تنسيق أعمال التنظيف وتجهيز الوحدة بما يضمن جاهزيتها للضيف."),
    ("الصيانة والمتابعة", "متابعة احتياجات الصيانة والتنسيق لمعالجة الملاحظات والمشكلات."),
    ("إدارة المنصات", "إدارة حسابات ومنصات الإيجار قصير المدى وتحسين ظهور الوحدة."),
    ("إدارة التقييمات", "متابعة تقييمات الضيوف وتحسين تجربة الإقامة وبناء سمعة أفضل للوحدة."),
    ("متابعة الأداء", "متابعة مؤشرات الأداء والإشغال والإيرادات وتقديم رؤية واضحة للمالك."),
]
for part in (0, 1):
    s = new_slide()
    slide_header(s, "OUR SERVICES", f"خدماتنا  {'١ / ٢' if part == 0 else '٢ / ٢'}")
    chunk = services[part*4:part*4+4]
    cw = Inches(5.75)
    xs = [SW - Inches(0.7) - cw, Inches(0.7)]  # RTL: العمود الأيمن أولًا
    for i, (t, d) in enumerate(chunk):
        col, row = i % 2, i // 2
        x = xs[col]
        y = Inches(2.15) + row * Inches(2.3)
        rect(s, x, y, cw, Inches(2.05),
             fill=WHITE, line_color=LINE, line_w=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        rect(s, x + cw - Inches(0.14), y + Inches(0.25), Pt(2.4), Inches(1.55), fill=GOLD)
        gold_num(s, x + cw - Inches(0.85), y + Inches(0.28), f"0{part*4 + i + 1}")
        _, tf = box(s, x + Inches(0.3), y + Inches(0.3), cw - Inches(1.45), Inches(0.5))
        para(tf, t, F_HEAD, 19, INK, bold=True, first=True)
        _, tf = box(s, x + Inches(0.3), y + Inches(0.95), cw - Inches(1.45), Inches(1.0))
        para(tf, d, F_BODY, 13, INK_SOFT, line=1.4, first=True)

# ================================================================ 7 · آلية العمل
s = new_slide()
slide_header(s, "HOW IT WORKS", "آلية العمل — من التعرّف إلى النتائج")
steps = ["نتعرف على وحدتك", "نجهّز خطة التشغيل", "نبدأ الإدارة والتشغيل",
         "نتابع الضيوف والحجوزات", "نراقب الأداء والتقييمات", "نرفع مستوى التجربة باستمرار"]
hline(s, Inches(0.9), Inches(3.06), Inches(11.5), GOLD_SOFT, 1.5)
w = Inches(1.95)
for i, st in enumerate(steps):
    x = Inches(12.4) - w - i * (w + Inches(0.05))
    dot = rect(s, x + w/2 - Inches(0.3), Inches(2.78), Inches(0.56), Inches(0.56),
               fill=INK, line_color=GOLD, line_w=1.5, shape=MSO_SHAPE.OVAL)
    tfd = dot.text_frame
    tfd.margin_left = tfd.margin_right = tfd.margin_top = tfd.margin_bottom = 0
    pd = tfd.paragraphs[0]; pd.alignment = PP_ALIGN.CENTER
    rd = pd.add_run(); rd.text = str(i+1)
    set_font(rd, F_LAT, 16, GOLD_SOFT)
    _, tf = box(s, x - Inches(0.12), Inches(3.62), w + Inches(0.24), Inches(1.4))
    para(tf, st, F_BODY, 14, INK, bold=True, align=PP_ALIGN.CENTER, line=1.35, first=True)
_, tf = box(s, Inches(0.7), Inches(5.75), Inches(12.0), Inches(0.6))
para(tf, "التقييم الأولي وخطة التشغيل — بدون أي رسوم مقدمة. أتعابنا مرتبطة بالنتائج الفعلية.",
     F_BODY, 14, GOLD_DEEP, align=PP_ALIGN.CENTER, first=True)

# ================================================================ 8 · لماذا رِفان
s = new_slide()
slide_header(s, "WHY RIFAN", "لماذا رِفان؟")
why = [
    ("إدارة متكاملة", "لا نترك المالك أمام تفاصيل التشغيل اليومية."),
    ("تجربة ضيافة احترافية", "نهتم بالتفاصيل التي تصنع الفرق في تجربة الضيف."),
    ("متابعة مستمرة", "نراقب التشغيل والحجوزات والتقييمات والأداء."),
    ("حل واحد للمالك", "من الحجوزات إلى التنظيف والصيانة وخدمة الضيوف."),
    ("شفافية في الأداء", "المالك على اطلاع واضح على أداء وحدته."),
    ("خبرة تشغيلية", "نعتمد على خبرة عملية وفريق متخصص في تشغيل الوحدات."),
]
for i, (t, d) in enumerate(why):
    col, row = i % 3, i // 3
    x = Inches(12.65) - Inches(3.95) - col * Inches(4.12)
    y = Inches(2.2) + row * Inches(2.25)
    rect(s, x, y, Inches(3.95), Inches(2.0), fill=WHITE, line_color=LINE, line_w=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    hline(s, x + Inches(3.35), y + Inches(0.35), Inches(0.35), GOLD, 3.5)
    _, tf = box(s, x + Inches(0.25), y + Inches(0.5), Inches(3.45), Inches(0.5))
    para(tf, t, F_HEAD, 17, INK, bold=True, first=True)
    _, tf = box(s, x + Inches(0.25), y + Inches(1.05), Inches(3.45), Inches(0.85))
    para(tf, d, F_BODY, 12.5, INK_SOFT, line=1.4, first=True)

# ================================================================ 9 · المنصات
s = new_slide(INK)
slide_header_dark = None
latin_label(s, Inches(0.7), Inches(0.52), Inches(12.0), "BOOKING PLATFORMS", GOLD_SOFT, 11)
_, tf = box(s, Inches(0.7), Inches(0.82), Inches(12.0), Inches(0.85))
para(tf, "نُدير وحدتك على كبرى منصات الحجز", F_HEAD, 32, IVORY, bold=True, first=True)
hline(s, Inches(11.1), Inches(1.62), Inches(1.6), GOLD, 2.0)
platforms = ["AIRBNB", "BOOKING.COM", "GATHERN", "EXPEDIA", "TRIP.COM"]
# صورة وحدة على اليسار بإطار ذهبي
pimg = crop("bedroom.jpg", "platform_panel.jpg", 4.4 / 4.45)
s.shapes.add_picture(pimg, Inches(0.7), Inches(2.15), Inches(4.4), Inches(4.45))
rect(s, Inches(0.7), Inches(2.15), Inches(4.4), Inches(4.45), fill=None, line_color=GOLD, line_w=1.2)
_, tf = box(s, Inches(5.6), Inches(2.15), Inches(7.05), Inches(1.1))
para(tf, "إدارة حسابات الوحدة، تحسين الإعلان والصور والتسعير، ومزامنة التقويمات — لتحقيق أعلى إشغال وأفضل عائد.",
     F_BODY, 15, IVORY_DEEP, line=1.5, first=True)
cw2 = Inches(6.85)
for i, p in enumerate(platforms):
    y = Inches(3.4) + i * Inches(0.66)
    c = rect(s, Inches(12.65) - cw2, y, cw2, Inches(0.58), fill=None, line_color=GOLD, line_w=1.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tfc = c.text_frame; tfc.word_wrap = False
    tfc.margin_left = tfc.margin_right = tfc.margin_top = tfc.margin_bottom = 0
    pc = tfc.paragraphs[0]; pc.alignment = PP_ALIGN.CENTER
    rc = pc.add_run(); rc.text = p
    set_font(rc, F_LAT, 13, GOLD_SOFT, spc=120)
_, tf = box(s, Inches(0.7), Inches(7.05), Inches(3), Inches(0.3))
para(tf, "RIFAN", F_LAT, 10, MUTED, align=PP_ALIGN.LEFT, rtl=False, spc=500, first=True)

# ================================================================ 10 · قصص نجاح
s = new_slide()
slide_header(s, "SUCCESS STORIES", "قصص نجاح — أرقام تتحدث")
stories = [
    ("+41%", "ارتفاع معدل الإشغال", "وحدة بغرفة واحدة انتقلت من إشغال متذبذب إلى إشغال مستقر خلال أول أربعة أشهر من التشغيل.", "living.jpg"),
    ("+38%", "نمو متوسط الإيراد الشهري", "إعادة بناء التسعير وتحسين الظهور على منصات الحجز رفع متوسط الإيراد الشهري للوحدة.", "bedroom.jpg"),
    ("4.9/5", "تقييم الضيوف", "متابعة احترافية لخدمة الضيوف والتقييمات حوّلت الوحدة إلى خيار موثوق ومفضّل.", "about.jpg"),
]
for i, (num, t, d, img) in enumerate(stories):
    x = Inches(12.65) - Inches(3.95) - i * Inches(4.12)
    y = Inches(2.2)
    rect(s, x, y, Inches(3.95), Inches(4.3), fill=WHITE, line_color=LINE, line_w=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    simg = crop(img, f"st{i}.jpg", 3.95 / 1.7)
    s.shapes.add_picture(simg, x, y, Inches(3.95), Inches(1.7))
    overlay(s, x, y, Inches(3.95), Inches(1.7), INK, 72)
    _, tf = box(s, x + Inches(0.2), y + Inches(0.42), Inches(3.55), Inches(0.9))
    para(tf, num, F_LAT, 40, GOLD_SOFT, bold=True, align=PP_ALIGN.CENTER, rtl=False, first=True)
    _, tf = box(s, x + Inches(0.25), y + Inches(1.95), Inches(3.45), Inches(0.5))
    para(tf, t, F_HEAD, 17, GOLD_DEEP, bold=True, first=True)
    hline(s, x + Inches(3.1), y + Inches(2.5), Inches(0.5), GOLD, 2.0)
    _, tf = box(s, x + Inches(0.25), y + Inches(2.7), Inches(3.45), Inches(1.4))
    para(tf, d, F_BODY, 12.5, INK_SOFT, line=1.45, first=True)

# ================================================================ 11 · الأسئلة الشائعة
s = new_slide()
slide_header(s, "FAQ", "أسئلة يسألها كل مالك")
faqs = [
    ("هل تبقى الوحدة مسجلة باسمي؟", "نعم، دائمًا. نعمل بعقد إدارة وتشغيل يوثّق ملكيتك الكاملة للوحدة، ونتولى فقط إدارتها وتشغيلها."),
    ("كيف أتابع أداء وحدتي؟", "تقارير دورية واضحة تشمل الإشغال والإيرادات وتقييمات الضيوف، مع تواصل مباشر مع فريقنا."),
    ("هل توجد رسوم قبل بدء التشغيل؟", "التقييم الأولي وخطة التشغيل بدون رسوم. أتعابنا مرتبطة بالنتائج الفعلية."),
    ("هل يمكنني إيقاف التعاون في أي وقت؟", "عقد مرن بفترة إشعار مسبق متفق عليها — بدون التزامات طويلة المدى."),
]
for i, (q, a) in enumerate(faqs):
    col, row = i % 2, i // 2
    x = Inches(12.65) - Inches(5.9) - col * Inches(6.12)
    y = Inches(2.15) + row * Inches(2.3)
    rect(s, x, y, Inches(5.9), Inches(2.05), fill=WHITE, line_color=LINE, line_w=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    _, tf = box(s, x + Inches(0.3), y + Inches(0.28), Inches(5.3), Inches(0.5))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    p._p.get_or_add_pPr().set('rtl', '1')
    r1 = p.add_run(); r1.text = "س:  "
    set_font(r1, F_HEAD, 15, GOLD_DEEP, bold=True)
    r2 = p.add_run(); r2.text = q
    set_font(r2, F_HEAD, 15, INK, bold=True)
    _, tf = box(s, x + Inches(0.3), y + Inches(0.95), Inches(5.3), Inches(1.0))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    p._p.get_or_add_pPr().set('rtl', '1')
    r3 = p.add_run(); r3.text = a
    set_font(r3, F_BODY, 12.5, INK_SOFT)
    tf.paragraphs[0].line_spacing = 1.4

# ================================================================ 12 · تواصل معنا (QR)
s = new_slide(INK)
corner_marks(s, Inches(0.55), Inches(0.55), SW - Inches(1.1), SH - Inches(1.1), GOLD, Inches(0.7), Pt(2.0))
latin_label(s, Inches(0.7), Inches(0.65), Inches(12.0), "CONTACT US", GOLD_SOFT, 11)
_, tf = box(s, Inches(0.7), Inches(0.95), Inches(12.0), Inches(0.85))
para(tf, "تواصل معنا — امسح الرمز", F_HEAD, 32, IVORY, bold=True, first=True)
hline(s, Inches(11.1), Inches(1.78), Inches(1.6), GOLD, 2.0)

qr_specs = [
    ("WHATSAPP", "+966 50 093 9971", f"https://wa.me/{WHATSAPP}"),
    ("INSTAGRAM", "I0.rifan", "https://instagram.com/I0.rifan"),
    ("TIKTOK", "@I0.rifan", "https://tiktok.com/@I0.rifan"),
]
w = Inches(3.3)
total = w * 3 + Inches(0.4) * 2
for i, (name, handle, url) in enumerate(qr_specs):
    x = SW/2 - total/2 + i * (w + Inches(0.4))
    y = Inches(2.35)
    card = rect(s, x, y, w, Inches(3.9), fill=IVORY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    make_qr(url, f"qr_{i}.png")
    s.shapes.add_picture(f"qr_{i}.png", x + w/2 - Inches(1.15), y + Inches(0.35), height=Inches(2.3))
    _, tf = box(s, x + Inches(0.2), y + Inches(2.8), w - Inches(0.4), Inches(0.4))
    para(tf, name, F_LAT, 14, GOLD_DEEP, bold=True, align=PP_ALIGN.CENTER, rtl=False, spc=250, first=True)
    _, tf = box(s, x + Inches(0.2), y + Inches(3.25), w - Inches(0.4), Inches(0.4))
    para(tf, handle, F_LAT, 13, INK_SOFT, align=PP_ALIGN.CENTER, rtl=False, first=True)
_, tf = box(s, Inches(1.5), Inches(6.55), Inches(10.333), Inches(0.45))
para(tf, "البريد الإلكتروني:  b.rifan.mk@gmail.com  ·  واتساب: \u200e+966 50 093 9971\u200e",
     F_BODY, 15, GOLD_SOFT, align=PP_ALIGN.CENTER, first=True)

# ================================================================ 13 · الختام
s = new_slide(INK)
closing = crop("cinematic.jpg", "closing_169.jpg", 13.333 / 7.5)
s.shapes.add_picture(closing, 0, 0, SW, SH)
overlay(s, 0, 0, SW, SH, INK, 85)
rect(s, Inches(0.32), Inches(0.32), SW - Inches(0.64), SH - Inches(0.64),
     fill=None, line_color=GOLD, line_w=1.0)
corner_marks(s, Inches(0.55), Inches(0.55), SW - Inches(1.1), SH - Inches(1.1), GOLD, Inches(0.7), Pt(2.0))
s.shapes.add_picture(LOGO, SW/2 - Inches(0.85), Inches(1.5), height=Inches(1.7))
_, tf = box(s, Inches(1.5), Inches(3.5), Inches(10.333), Inches(1.0))
para(tf, "وحدتك تستحق إدارة بمستوى ضيافة فاخر", F_HEAD, 34, IVORY, bold=True, align=PP_ALIGN.CENTER, first=True)
hline(s, SW/2 - Inches(1.1), Inches(4.5), Inches(2.2), GOLD, 1.5)
_, tf = box(s, Inches(1.5), Inches(4.75), Inches(10.333), Inches(0.5))
para(tf, "Luxury Hospitality + Professional Property Management", F_LAT, 13, GOLD_SOFT,
     align=PP_ALIGN.CENTER, rtl=False, spc=350, first=True)
_, tf = box(s, Inches(1.5), Inches(5.6), Inches(10.333), Inches(0.5))
para(tf, "رِفان  ·  إدارة وتشغيل الوحدات السكنية للإيجار قصير المدى", F_BODY, 14, MUTED,
     align=PP_ALIGN.CENTER, first=True)
_, tf = box(s, Inches(1.5), Inches(6.2), Inches(10.333), Inches(0.5))
para(tf, "\u200e+966 50 093 9971\u200e  (واتساب)   ·   \u200eb.rifan.mk@gmail.com\u200e   ·   \u200einstagram.com/I0.rifan\u200e   ·   \u200etiktok.com/@I0.rifan\u200e",
     F_BODY, 12, GOLD_SOFT, align=PP_ALIGN.CENTER, first=True)

OUT = r"C:\Projects\rivan\RIFAN_Services_Presentation.pptx"
prs.save(OUT)
import os
for i in range(3):
    os.remove(f"qr_{i}.png")
print("Saved:", OUT, f"({os.path.getsize(OUT)//1024} KB, {len(prs.slides._sldIdLst)} slides)")
